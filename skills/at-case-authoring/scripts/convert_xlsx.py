#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""convert_xlsx.py — xlsx「用例」sheet → cases_standard.yaml（原始描述全保留 + 切分 + element-map 初稿）。

三步流程：
  1. 机械转换：提取 xlsx 原始字段到 raw_* 字段，逐字保留；
     同时按规范做机械预处理（拆行、去行号、提模块）到规范字段。
  2. 切分：按模块分组 + token 预算切分为 slices/（供 AI 初步生成逐片处理）。
  3. element-map 初稿：机械提取步骤中的 UI 目标（ui_name），id_name/role 留 TBD。

产出到输出目录：
  cases_standard.yaml           # 全量（含 raw_*，供人工校对 / 校验）
  element-map.yaml              # 初稿（ui_name 已提取，id_name/role=TBD 待开发填）
  slices/<module>_<seq>.yaml    # 按模块 + token 预算切分（AI 初步生成逐片输入）

用法:
    convert_xlsx.py --input 用例.xlsx --output <目录> [--app deepin-reader] [--budget 16000]
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

STEP_NUM = re.compile(r"^\s*\d+[\.、．)）]?\s*")
_CJK_RE = re.compile(r"[\u4e00-\u9fff]")
_ACTION_WIDGET = re.compile(r"(?:点击|单击|双击|右键|左键|选择|选中|勾选|点选|按下|切换到|打开|关闭)(?:了|一下)?([\u4e00-\u9fffA-Za-z0-9]{1,6}?(?:按钮|输入框|搜索框|下拉框|菜单项|图标|滑块|复选框|标签页|菜单))")
_ACTION_CTX = re.compile(r"(?:点击|单击|双击|右键|左键|选择|选中|勾选|打开|关闭|切换到|进入|点选|按下)\s*[\u4e00-\u9fff]{0,6}?[\x22\x27\u201c\u201d]([^\x22\x27\u201c\u201d]{1,24})[\x22\x27\u201c\u201d]")
_INPUT_CTX = re.compile(r"(?:输入|键入|粘贴|填写|内容为|输入内容)\s*[\u4e00-\u9fff]{0,6}?[\x22\x27\u201c\u201d]([^\x22\x27\u201c\u201d]{1,24})[\x22\x27\u201c\u201d]")


def split_lines(text: str) -> list[str]:
    if not text:
        return []
    out = []
    for line in str(text).replace("\r", "").split("\n"):
        line = line.strip()
        if line:
            out.append(STEP_NUM.sub("", line))
    return out


def module_short(module: str) -> str:
    """/path/to/查找(#116101) → 查找"""
    if not module:
        return "未分类"
    seg = module.rstrip("/").split("/")[-1]
    return re.sub(r"\(#.*\)\s*$", "", seg).strip() or "未分类"


def estimate_tokens(text: str) -> int:
    """与 at-suite-generator 一致的 token 估算：CJK~0.7 tok/char，其他~0.3。"""
    if not text:
        return 0
    cjk = len(_CJK_RE.findall(text))
    other = len(text) - cjk
    return int(cjk * 0.7 + other * 0.3)


def _case_tokens(case: dict) -> int:
    return estimate_tokens(
        " ".join(str(case.get(k) or "") for k in ("id", "title", "module", "precondition", "steps", "expected"))
    )


def _slugify(text: str) -> str:
    cleaned = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", text).strip("_")
    return cleaned[:48] or "misc"


def extract_ui_targets(cases: list[dict]) -> list[str]:
    """机械提取 UI 目标初稿：动作动词后的引号文本 + 明显控件词。去重保序。

    仅候选初稿：先收集输入数据（输入/键入/粘贴 后引号内容）作排除集，
    动作上下文引号文本和控件词短语不入排除集才保留。
    最终由 AI 初步生成阶段精修 + 测试人员校对。
    """
    input_data: set[str] = set()
    for c in cases:
        for step in c.get("steps") or []:
            s = str(step)
            for m in _INPUT_CTX.finditer(s):
                input_data.add(m.group(1).strip())

    seen: list[str] = []
    for c in cases:
        for step in c.get("steps") or []:
            s = str(step)
            for m in _ACTION_CTX.finditer(s):
                t = m.group(1).strip()
                if t and t not in input_data and t not in seen and any("\u4e00" <= ch <= "\u9fff" for ch in t):
                    seen.append(t)
            for m in _ACTION_WIDGET.finditer(s):
                w = m.group(1).strip()
                if w and w not in input_data and w not in seen and any("\u4e00" <= ch <= "\u9fff" for ch in w):
                    seen.append(w)
    return seen


def read_xlsx(path: Path, app: str, out_dir: Path, budget: int) -> None:
    try:
        import openpyxl
    except ImportError:
        print("需要 openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(2)
    try:
        import yaml
    except ImportError:
        print("需要 PyYAML: pip install pyyaml", file=sys.stderr)
        sys.exit(2)

    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    header = [str(c).strip() if c else "" for c in next(ws.iter_rows(values_only=True))]
    idx = {h: i for i, h in enumerate(header)}

    def col(row, *names):
        for n in names:
            j = idx.get(n)
            if j is not None and j < len(row) and row[j] is not None:
                return str(row[j])
        return ""

    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None for c in row):
            continue
        raw_pre = col(row, "前置条件", "前提条件", "预置条件")
        raw_steps = col(row, "步骤", "测试步骤", "操作步骤")
        raw_exp = col(row, "预期", "预期结果", "期望结果")
        cases.append({
            "id": str(col(row, "用例编号", "ID", "编号", "序号")),
            # 原始描述（逐字保留，测试人员校对对照用）
            "raw_title": col(row, "用例标题", "标题", "用例名称"),
            "raw_precondition": raw_pre,
            "raw_steps": raw_steps,
            "raw_expected": raw_exp,
            # 机械预处理字段（AI 在此基础上初步生成）
            "title": col(row, "用例标题", "标题", "用例名称"),
            "module": module_short(col(row, "所属模块", "模块", "功能模块")),
            "priority": col(row, "用例级别", "优先级", "级别"),
            "precondition": split_lines(raw_pre),
            "steps": split_lines(raw_steps),
            "expected": split_lines(raw_exp),
            "case_type": col(row, "用例类型", "类型", "测试类型"),
            "manual": False,
        })

    if not cases:
        print("xlsx 无有效用例行", file=sys.stderr)
        sys.exit(1)

    out_dir.mkdir(parents=True, exist_ok=True)

    # ── 1. 全量 cases_standard.yaml ──
    doc = {
        "app": app,
        "source": path.name,
        "converted_at": date.today().isoformat(),
        "case_count": len(cases),
        "generated_by": "convert_xlsx.py (mechanical)",
        "cases": cases,
    }
    full_path = out_dir / "cases_standard.yaml"
    full_path.write_text(yaml.safe_dump(doc, allow_unicode=True, sort_keys=False), encoding="utf-8")

    # ── 2. 切分：按模块分组 + token 预算 ──
    groups: dict[str, list[dict]] = {}
    order: list[str] = []
    for c in cases:
        mod = c.get("module") or "未分类"
        if mod not in groups:
            groups[mod] = []
            order.append(mod)
        groups[mod].append(c)

    slices_dir = out_dir / "slices"
    slices_dir.mkdir(exist_ok=True)
    slice_meta: list[dict] = []
    for mod in order:
        mod_cases = groups[mod]
        # 预算内切分（单条超预算也独立成片，不丢弃）
        seq = 1
        current: list[dict] = []
        cur_tokens = 0
        for case in mod_cases:
            t = _case_tokens(case)
            if current and cur_tokens + t > budget:
                slug = _slugify(mod)
                f = slices_dir / f"{slug}_{seq:03d}.yaml"
                f.write_text(yaml.safe_dump({
                    "app": app,
                    "module": mod,
                    "seq": seq,
                    "case_count": len(current),
                    "cases": current,
                }, allow_unicode=True, sort_keys=False), encoding="utf-8")
                slice_meta.append({"module": mod, "seq": seq, "file": f.name, "case_count": len(current)})
                current = []
                cur_tokens = 0
                seq += 1
            current.append(case)
            cur_tokens += t
        if current:
            slug = _slugify(mod)
            f = slices_dir / f"{slug}_{seq:03d}.yaml"
            f.write_text(yaml.safe_dump({
                "app": app,
                "module": mod,
                "seq": seq,
                "case_count": len(current),
                "cases": current,
            }, allow_unicode=True, sort_keys=False), encoding="utf-8")
            slice_meta.append({"module": mod, "seq": seq, "file": f.name, "case_count": len(current)})

    # ── 3. element-map.yaml 初稿 ──
    ui_targets = extract_ui_targets(cases)
    elem_doc = {
        "app": app,
        "version": "0.1",
        "updated": date.today().isoformat(),
        "note": "初稿：ui_name 由脚本机械提取，id_name/role=TBD 待开发人员填充",
        "elements": [
            {"desc": "TBD", "ui_name": u, "id_name": "TBD", "role": "TBD"}
            for u in ui_targets
        ],
    }
    elem_path = out_dir / "element-map.yaml"
    elem_path.write_text(yaml.safe_dump(elem_doc, allow_unicode=True, sort_keys=False), encoding="utf-8")

    print(f"转换完成: {len(cases)} 条 → {out_dir}")
    print(f"  全量: cases_standard.yaml ({len(cases)} 条)")
    print(f"  分片: slices/ ({len(slice_meta)} 片, 预算 {budget} tok)")
    print(f"  element-map 初稿: element-map.yaml ({len(ui_targets)} 个 UI 目标)")
    print("下一步: 按 SKILL.md 模式 B 做 AI 初步生成（逐片规范化 + 校对 element-map）")


def main() -> None:
    ap = argparse.ArgumentParser(description="xlsx 用例 → 标准 YAML（保留原始描述 + 切分 + element-map 初稿）")
    ap.add_argument("--input", required=True, help="xlsx 用例文件")
    ap.add_argument("--output", required=True, help="输出目录（生成 cases_standard.yaml + slices/ + element-map.yaml）")
    ap.add_argument("--app", default="", help="应用名")
    ap.add_argument("--budget", type=int, default=16000, help="每片 token 预算（默认 16000，与 at-suite-generator 一致）")
    args = ap.parse_args()

    in_path = Path(args.input)
    if not in_path.exists():
        print(f"文件不存在: {in_path}", file=sys.stderr)
        sys.exit(1)
    read_xlsx(in_path, args.app, Path(args.output), args.budget)


if __name__ == "__main__":
    main()
