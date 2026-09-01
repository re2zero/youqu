#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""validate_cases.py — 合规检查：校验用例是否符合《AT用例编写规范》。

输入：cases_standard.yaml（或 xlsx）+ 可选 element-map.yaml
输出：合规报告（stdout），exit code 0=无 error，1=有 error。

规则来源：references/spec.md（断言关键词 / 描述词 / 角色词 / 菜单词等）。
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

# ─── 关键词白名单（与规范一致）──────────────────────────────
ASSERT_KEYWORDS = ["检查", "确认", "查看", "验证", "是否", "应该",
                   "符合", "出现", "消失", "正确", "可见", "观察", "对比"]
DESC_KEYWORDS = ["显示", "被清空", "可以重新", "正常", "异常", "检查",
                 "查看", "应当", "应该"]
VAGUE_INPUT = re.compile(r"输入(任意|很|较|超|非常)?长?的?(字符|内容|文本|信息|长度)")
VAGUE_POS = re.compile(r"右上角|右下角|左上角|左下角|左侧|右侧|左边|右边|空白处|其他地方|其它地方|任意位置")
VAGUE_ACTION_START = re.compile(r"^(符合|没有|显示|出现|无|不出现|当)")
ACTION_VERBS = ["点击", "单击", "双击", "右键", "左键", "长按", "输入", "键入",
                "粘贴", "按", "快捷键", "回车", "滚轮", "拖拽", "打开", "关闭",
                "新建", "删除", "重命名", "清空", "展开", "收起", "切换",
                "勾选", "取消勾选", "等待", "选择", "选中", "进入", "退出", "执行"]
UNSUPPORTED_TARGETS = [
    ("颜色/高亮/置灰", ["背景变模糊", "语法高亮", "按钮置灰", "置灰", "背景模糊"]),
    ("光标/选区形态", ["光标形态", "选区形态", "光标位置"]),
    ("tooltip/悬停", ["Tip框", "悬停", "tooltip", "提示框"]),
    ("系统环境", ["dock", "Dock", "休眠", "待机", "锁屏", "开机自启", "桌面"]),
    ("外部硬件", ["触控板", "触摸屏", "麦克风", "音频", "手势"]),
    ("文件/网络依赖", ["U盘", "光驱", "smb", "SMB", "ftp", "FTP", "sftp", "网络", "挂载", "1G", "大文件"]),
    ("性能/资源", ["内存", "冷热启动", "长时间运行", "卡顿", "性能", "耗时"]),
    ("跨应用/破坏性", ["关机", "重启", "注销", "强制退出", "崩溃"]),
]



class Report:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, case_id: str, msg: str) -> None:
        self.errors.append(f"[{case_id}] {msg}")

    def warning(self, case_id: str, msg: str) -> None:
        self.warnings.append(f"[{case_id}] {msg}")

    def render(self) -> str:
        lines = []
        lines.append(f"errors: {len(self.errors)}  warnings: {len(self.warnings)}")
        for e in self.errors:
            lines.append(f"ERROR  {e}")
        for w in self.warnings:
            lines.append(f"WARN   {w}")
        return "\n".join(lines)


def strip_step(step: str) -> str:
    return re.sub(r"^\s*\d+[\.、．)）]?\s*", "", step).strip()


def validate_cases(cases: list[dict], report: Report) -> None:
    for case in cases:
        cid = str(case.get("id", "?"))
        title = str(case.get("title", ""))
        precondition = case.get("precondition") or []
        if isinstance(precondition, str):
            precondition = [precondition]
        steps = case.get("steps") or []
        if isinstance(steps, str):
            steps = [steps]
        expected = case.get("expected") or []
        if isinstance(expected, str):
            expected = [expected]

        # 1. 标题
        if not re.match(r"^【[^】]+】", title) and not case.get("manual"):
            report.warning(cid, f"标题缺【模块】前缀: {title[:30]}")


        # 2. 前置条件：禁止动作句
        for p in precondition:
            p = strip_step(p)
            if any(p.startswith(v) for v in ACTION_VERBS):
                report.warning(cid, f"前置条件含动作(应为状态): {p[:40]}")

        # 3. 步骤
        for step in steps:
            s = strip_step(step)
            if not s:
                continue
            if any(kw in s for kw in ASSERT_KEYWORDS):
                report.warning(cid, f"步骤含断言词(移到预期列): {s[:40]}")
            # 分号多动作检测：排除代码片段（含括号/引号/等号，或无中文的纯代码行）
            has_cjk = bool(re.search(r"[\u4e00-\u9fff]", s))
            is_code = bool(re.search(r"[\(\)\"\'=]", s)) or not has_cjk
            if ("；" in s or ";" in s) and not is_code:
                report.error(cid, f"步骤含分号多动作(应拆分): {s[:40]}")
            if VAGUE_ACTION_START.match(s) and not any(v in s for v in ACTION_VERBS):
                report.error(cid, f"步骤以状态描述开头(非操作): {s[:40]}")
            if VAGUE_POS.search(s):
                report.warning(cid, f"步骤位置无UI文本锚点: {s[:40]}")
            if VAGUE_INPUT.search(s):
                # 排除有具体 payload 的形式：冒号后内容 或 引号包裹内容
                m = re.search(r"[输键贴粘]入[^:：\n]{0,10}[:：]\s*(\S+)", s)
                q = re.search(r"[输键贴粘]入[\u4e00-\u9fff]{0,6}[\x22\x27\u201c\u201d]([^\x22\x27\u201c\u201d]{1,20})[\x22\x27\u201c\u201d]", s)
                has_payload = bool(
                    (m and len(m.group(1)) >= 2 and not any(k in m.group(1) for k in DESC_KEYWORDS))
                    or (q and not any(k in q.group(1) for k in DESC_KEYWORDS))
                )
                if not has_payload:
                    report.error(cid, f"输入无具体数据: {s[:40]}")
            if re.search(r"[输键贴粘]入\s*[:：]\s*[^ ]*(显示|被清空|可以重新|正常|异常|检查|查看|应当|应该)",
                         s):
                report.error(cid, f"输入内容含描述词: {s[:40]}")

        # 4. 预期：禁止"操作正常/无异常"；检测不可自动化目标
        for e in expected:
            e = strip_step(e)
            if not e:
                continue
            if re.search(r"操作正常|无异常|功能正常|显示正常|一切正常", e):
                report.warning(cid, f"预期过泛(应可断言): {e[:40]}")
        # 不可自动化标记一致性：manual:true 必须标题带【人工】，反之亦然
        is_manual = bool(case.get("manual"))
        has_label = title.startswith("【人工】")
        if is_manual != has_label:
            report.error(
                cid,
                f"manual({is_manual}) 与标题【人工】标记不一致"
                + ("，manual:true 须标题加【人工】" if is_manual else "，标题标【人工】须 manual:true"),
            )
        if is_manual and not (case.get("reason") or "").strip():
            report.error(cid, "manual:true 用例须填 reason 注明不可自动化类别（见 spec §7）")
        if not is_manual:
            joined = " ".join(steps + expected)
            for cat, kws in UNSUPPORTED_TARGETS:
                if any(k in joined for k in kws):
                    report.warning(cid, f"断言目标疑似不可自动化({cat})，应标【人工】+manual:true")
def validate_element_map(cases: list[dict], elem_map: dict | None, report: Report) -> None:
    if not elem_map:
        report.warning("?", "无 element-map.yaml，无法校验 UI 目标映射")
        return
    ui_names = {e.get("ui_name") for e in elem_map.get("elements", [])}
    for case in cases:
        cid = str(case.get("id", "?"))
        for step in (case.get("steps") or []):
            s = strip_step(step)
            if not s:
                continue
            # 提取带引号的控件文本（\x22=双引号 \x27=单引号 \u201c\u201d=弯引号）
            quoted = re.findall(r"[\x22\x27\u201c\u201d]([^\x22\x27\u201c\u201d ]{1,20})[\x22\x27\u201c\u201d]", s)
            for q in quoted:
                if q not in ui_names:
                    report.warning(cid, f"UI目标不在element-map中: {q}")


def load_yaml(path: Path) -> dict | None:
    try:
        import yaml
    except ImportError:
        print("需要 PyYAML: pip install pyyaml", file=sys.stderr)
        sys.exit(2)
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def main() -> None:
    ap = argparse.ArgumentParser(description="用例合规检查")
    ap.add_argument("input", help="cases_standard.yaml 或 xlsx")
    ap.add_argument("--element-map", help="element-map.yaml（可选）")
    args = ap.parse_args()

    in_path = Path(args.input)
    report = Report()

    if in_path.suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            print("需要 openpyxl: pip install openpyxl", file=sys.stderr)
            sys.exit(2)
        wb = openpyxl.load_workbook(in_path, read_only=True)
        ws = wb.worksheets[0]
        header = [str(c).strip() if c else "" for c in next(ws.iter_rows(values_only=True))]
        idx = {h: i for i, h in enumerate(header)}
        cases = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(c is None for c in row):
                continue
            def col(name):
                j = idx.get(name)
                return row[j] if j is not None and j < len(row) and row[j] is not None else ""
            cases.append({
                "id": str(col("用例编号")) or "?",
                "title": col("用例标题"),
                "precondition": col("前置条件"),
                "steps": col("步骤"),
                "expected": col("预期"),
            })
        print(f"loaded {len(cases)} cases from xlsx")
    else:
        data = load_yaml(in_path)
        if not data or "cases" not in data:
            print("yaml 需含 cases 键", file=sys.stderr)
            sys.exit(1)
        cases = data["cases"]
        print(f"loaded {len(cases)} cases from yaml")

    elem_map = load_yaml(Path(args.element_map)) if args.element_map else None

    validate_cases(cases, report)
    validate_element_map(cases, elem_map, report)

    print(report.render())
    sys.exit(0 if not report.errors else 1)


if __name__ == "__main__":
    main()
