# SPDX-FileCopyrightText: 2026 UnionTech Software Technology Co., Ltd.
#
# SPDX-License-Identifier: GPL-2.0-only

"""Parse xlsx/csv test case files into raw cases.yaml (format conversion only).

Semantic mapping (action/key/text/element_ref/selector) is done by the AI
in the session, not by this CLI command. Use ``youqu at tree-info`` to
produce a compact at-tree listing for AI context.

Pipeline:
1. Read xlsx/csv -> normalized case list
2. Build CasesDoc with step_type classification (keyword heuristic for steps
   column; assert steps from expected column). AI refines in Step 3.
3. Write raw cases.yaml
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

_ASSERT_KEYWORDS = frozenset(
    {
        "检查",
        "确认",
        "查看",
        "验证",
        "是否",
        "应该",
        "符合",
        "出现",
        "消失",
        "正确",
        "可见",
    }
)


def _guess_step_type(description: str) -> "StepType":  # noqa: F821
    from src.at.parser.models import StepType

    if any(kw in description for kw in _ASSERT_KEYWORDS):
        return StepType.assert_
    return StepType.action


COLUMN_ALIASES = {
    "id": ["用例编号", "ID", "编号", "序号"],
    "title": ["用例标题", "标题", "用例名称", "测试点"],
    "module": ["所属模块", "模块", "功能模块", "测试模块"],
    "priority": ["用例级别", "优先级", "级别", "重要程度"],
    "precondition": ["前置条件", "前提条件", "预置条件"],
    "steps": ["步骤", "测试步骤", "操作步骤", "用例步骤"],
    "expected": ["预期", "预期结果", "期望结果", "预期输出"],
    "case_type": ["用例类型", "类型", "测试类型"],
}


def find_column(record: dict, aliases: list[str]) -> str:
    for alias in aliases:
        if alias in record:
            val = record[alias]
            if val:
                return val
    for key in record:
        if any(a.lower() in key.lower() for a in aliases):
            return record[key]
    return ""


def read_xlsx(filepath: str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Error: openpyxl not installed. pip install openpyxl")
        sys.exit(1)
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else "" for h in rows[0]]
    data = []
    for row_idx, row in enumerate(rows[1:], start=2):
        record = {"source_row": row_idx}
        for col_idx, value in enumerate(row):
            if col_idx < len(headers) and headers[col_idx]:
                record[headers[col_idx]] = str(value).strip() if value is not None else ""
        data.append(record)
    return data


def read_csv_file(filepath: str) -> list[dict]:
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        data = []
        for row_idx, row in enumerate(reader, start=2):
            record = {"source_row": row_idx}
            record.update({k.strip(): v.strip() if v else "" for k, v in row.items() if k})
            data.append(record)
        return data


def normalize_cases(raw_data: list[dict]) -> list[dict]:
    cases = []
    for idx, record in enumerate(raw_data):
        cases.append(
            {
                "id": find_column(record, COLUMN_ALIASES["id"]) or str(idx).zfill(3),
                "title": find_column(record, COLUMN_ALIASES["title"]),
                "module": find_column(record, COLUMN_ALIASES["module"]),
                "priority": find_column(record, COLUMN_ALIASES["priority"]),
                "precondition": find_column(record, COLUMN_ALIASES["precondition"]),
                "steps": find_column(record, COLUMN_ALIASES["steps"]),
                "expected": find_column(record, COLUMN_ALIASES["expected"]),
                "case_type": find_column(record, COLUMN_ALIASES["case_type"]),
                "source_row": record.get("source_row", 0),
            }
        )
    return cases


def _compact_at_tree(at_tree_text: str) -> str:
    try:
        import yaml

        tree = yaml.safe_load(at_tree_text)
    except Exception:
        return at_tree_text

    if not tree or not isinstance(tree, dict):
        return at_tree_text

    lines: list[str] = []

    def _walk(nodes: list[dict], parent_path: list[str]) -> None:
        for node in nodes:
            if not isinstance(node, dict):
                continue
            nid = node.get("id", "")
            role = node.get("role", "")
            name = node.get("name", "")
            obj_name = node.get("object_name", "")
            path = parent_path + [nid] if nid else parent_path
            parent_str = " > ".join(path[:-1]) if len(path) > 1 else ""
            parts = ["{} | role: {} | name: {} | object_name: {}".format(nid, role, name, obj_name)]
            if parent_str:
                parts.append(" | parent: {}".format(parent_str))
            lines.append("".join(parts))
            children = node.get("children", [])
            if children:
                _walk(children, path)

    tree_nodes = tree.get("tree", [])
    if isinstance(tree_nodes, list):
        _walk(tree_nodes, [])

    return "\n".join(lines)


_TREE_INFO_KEEP_FIELDS = (
    "id",
    "role",
    "name",
    "object_name",
    "accessible_id",
    "classification",
    "comment",
    "annotation_status",
    "children",
)


def _simplify_tree(nodes: list[dict]) -> list[dict]:
    result = []
    for node in nodes:
        if not isinstance(node, dict):
            continue
        simplified = {k: node.get(k, "") for k in _TREE_INFO_KEEP_FIELDS if k != "children"}
        children = node.get("children", [])
        if children:
            simplified["children"] = _simplify_tree(children)
        result.append(simplified)
    return result


def compact_at_tree_to_file(at_tree_path: str, output_path: str, fmt: str = "yaml") -> None:
    raw = Path(at_tree_path).read_text(encoding="utf-8")
    if fmt == "text":
        compacted = _compact_at_tree(raw)
        Path(output_path).write_text(compacted, encoding="utf-8")
        print(
            "Wrote {} ({} bytes, {} nodes)".format(
                output_path, len(compacted), len(compacted.splitlines())
            )
        )
        return

    try:
        import yaml

        tree = yaml.safe_load(raw)
    except Exception:
        Path(output_path).write_text(raw, encoding="utf-8")
        return

    if not tree or not isinstance(tree, dict):
        Path(output_path).write_text(raw, encoding="utf-8")
        return

    tree_nodes = tree.get("tree", [])
    if not isinstance(tree_nodes, list):
        tree_nodes = []

    try:
        from src.at.scanner.merger import classify_nodes

        classify_nodes(tree_nodes)
    except ImportError:
        pass

    simplified = _simplify_tree(tree_nodes)
    output_data = {"version": "1.0", "tree": simplified}

    content = yaml.dump(
        output_data,
        allow_unicode=True,
        default_flow_style=False,
        sort_keys=False,
    )
    Path(output_path).write_text(content, encoding="utf-8")

    node_count = sum(1 for _ in _iter_nodes(simplified))
    print("Wrote {} ({} bytes, {} nodes)".format(output_path, len(content), node_count))


def _iter_nodes(nodes: list[dict]):
    for node in nodes:
        yield node
        for child in _iter_nodes(node.get("children", [])):
            yield child


def parse_to_cases(input_path: str, output_path: str, at_tree_path: str = "") -> None:
    """Parse xlsx/csv into raw cases.yaml (format conversion only).

    Semantic mapping (action/key/text/element_ref/selector) is done by the AI
    in the session, not by this CLI command.
    """
    path = Path(input_path)
    if not path.exists():
        print("Error: input file not found: {}".format(input_path))
        sys.exit(1)

    ext = path.suffix.lower()
    if ext in (".xlsx", ".xls"):
        raw_data = read_xlsx(input_path)
    elif ext == ".csv":
        raw_data = read_csv_file(input_path)
    else:
        print("Error: unsupported format: {}. Use .xlsx or .csv".format(ext))
        sys.exit(1)

    if not raw_data:
        print("Error: no data rows found in {}".format(input_path))
        sys.exit(1)

    cases = normalize_cases(raw_data)
    print("Read {} cases from {}".format(len(cases), input_path))

    if at_tree_path:
        print(
            "Warning: --at-tree is deprecated for parse. Semantic mapping is now done by AI. Use 'youqu at tree-info' instead."
        )

    from src.at.parser.models import CaseStep, CaseSuite, CasesDoc, CasesMetadata, StepType

    suites = []
    for case in cases:
        steps = []
        step_texts = case.get("steps", "")
        for line in step_texts.split("\n"):
            line = line.strip()
            if not line:
                continue
            steps.append(
                CaseStep(
                    step_type=_guess_step_type(line),
                    description=line,
                )
            )
        expected_texts = case.get("expected", "")
        for line in expected_texts.split("\n"):
            line = line.strip()
            if not line:
                continue
            steps.append(
                CaseStep(
                    step_type=StepType.assert_,
                    description=line,
                )
            )
        if steps:
            suites.append(
                CaseSuite(
                    id="case_{}".format(case["id"]),
                    name=case.get("title", ""),
                    module=case.get("module", ""),
                    description=case.get("title", ""),
                    status="active",
                    steps=steps,
                )
            )

    doc = CasesDoc(
        metadata=CasesMetadata(source=path.name),
        cases=suites,
    )

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    try:
        import yaml

        content = yaml.dump(
            doc.model_dump(mode="json", by_alias=False, exclude_none=True),
            allow_unicode=True,
            default_flow_style=False,
            sort_keys=False,
        )
    except ImportError:
        import json

        content = json.dumps(
            doc.model_dump(mode="json", by_alias=False, exclude_none=True),
            ensure_ascii=False,
            indent=2,
        )

    out.write_text(content, encoding="utf-8")
    total_steps = sum(len(s.steps) for s in doc.cases)
    print("Wrote {} ({} suites, {} raw steps)".format(output_path, len(doc.cases), total_steps))
